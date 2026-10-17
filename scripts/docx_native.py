"""docx для носителя языка: «Русский | Узбекский (перевод ИИ) | Правки» по всему тексту лендинга.

Русский берётся из собранной страницы ru/index.html, узбекский — из src/i18n/uz.json, порядок — как на странице.
Плюс скрытый текст (вкладка, превью ссылки в Telegram, подписи картинок) и политика конфиденциальности.
Формат тот же, что у бонуса (research/19-bonus-10-kanalov/build.py): альбомный лист, правая колонка пустая.

Запуск из корня репозитория после сборки: python3 scripts/docx_native.py
"""
from __future__ import annotations

import html
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
UZ = json.loads((ROOT / "src" / "i18n" / "uz.json").read_text(encoding="utf-8"))
OUT_DIR = ROOT.parent.parent / "research" / "13-lending-intensiva" / "для-носителя"
DOCX_NAME = "Лендинг мини-курса RU-UZ — на вычитку.docx"
DESK = Path.home() / "Desktop"
LINE = " / "  # разделитель строк одного заголовка в таблице

SECTIONS = {
    "b1": "1. Первый экран и блок «На мини-курсе»",
    "cta": "Кнопка (стоит на странице 5 раз)",
    "b2": "2. Уроки мини-курса и сертификат",
    "b3": "3. Кто ведёт",
    "b4": "4. Призыв и частые вопросы",
    "ft": "Подвал страницы",
    "hook": "Заголовки первого экрана под рекламу (меняются по объявлению, из которого пришёл человек)",
}
# заголовки из двух-трёх строк — одной строкой таблицы
GROUPS = [("b1.h1a", "b1.h1b", "b1.lead"), ("b2.h2a", "b2.h2b"), ("b3.h2a", "b3.h2b"), ("b4.h2a", "b4.h2b"), ("b4.faq_a", "b4.faq_b")]
LABELS = {
    "b1.kick": "надпись над заголовком", "b1.h1a": "главный заголовок и подзаголовок", "cta.start": "текст кнопки",
    "b1.learn_k": "подпись списка", "b1.bonus_k": "подпись", "b1.bonus_t": "название бонуса", "b1.bonus_d": "под бонусом",
    "b1.file_end": "последняя страница превью бонуса",
    "b2.h2a": "заголовок блока", "b2.s1": "надпись на кадре урока 1", "b2.s1a": "подпись картинки", "b2.s1b": "подпись картинки",
    "b2.s1c": "плашка на кадре", "b2.s2": "надпись на кадре урока 2", "b2.s3": "надпись на кадре урока 3",
    "b2.s3a": "плашка на кадре", "b2.s4live": "плашка на кадре", "b2.s4": "надпись на кадре практики",
    "b2.cd_k": "сертификат", "b2.cd_sub": "сертификат", "b2.cd_name": "сертификат: место для имени", "b2.cd_what": "сертификат",
    "b2.cd_sign": "сертификат: подпись", "b2.cd_role": "сертификат: под подписью", "b2.cert": "рядом с сертификатом",
    "b3.h2a": "заголовок блока", "b3.name": "имя", "b3.role": "должность", "b4.h2a": "заголовок блока",
    "b4.bonus_k": "подпись", "b4.bonus_t": "название бонуса", "b4.under": "под кнопкой", "b4.faq_a": "заголовок",
    "ft.entity": "реквизиты", "ft.bank": "реквизиты банка", "ft.privacy": "ссылка", "ft.disc": "оговорка",
}
INTRO = [
    "Это страница бесплатного мини-курса «ИИ на заказ» (по-узбекски — «AI buyurtma»): https://aibuyurtma.com. "
    "Сюда приходят из рекламы в Instagram и Facebook ребята 18–25 лет из Узбекистана, кнопка ведёт в Telegram-бота с уроками.",
    "Слева — русский текст. В середине — узбекский перевод, сделанный с помощью ИИ. Справа — пустая колонка для ваших правок.",
    "Строки идут в том же порядке, что и на странице, сверху вниз. В конце — скрытый текст (название вкладки, превью ссылки "
    "в Telegram, подписи к картинкам) и политика конфиденциальности.",
]
ASKS = [
    "Если узбекский звучит естественно и смысл совпадает с русским — напишите в правой колонке «ок».",
    "Если что-то не так — напишите в правой колонке весь правильный текст и выделите жёлтым то, что поменяли.",
    "Латиница; o‘ и g‘ — с апострофом ‘. К читателю — на «siz», как в рекламе (по-русски на странице «ты»).",
    f"Знак «{LINE.strip()}» в строке — это перенос строки в заголовке. Оставьте его там, где удобно разбить фразу.",
    "Цифры, суммы ($300, $1000, 2,5–5 mln so‘m) и порядок строк не меняем. Если число кажется странным — напишите об этом.",
    "Не добавляйте обещаний вроде kafolat, tez boyish, oson pul, passiv daromad — реклама Meta за такое блокируется.",
]
QUESTIONS = [
    "Название «AI buyurtma» (логотип и адрес сайта aibuyurtma.com): звучит нормально как название курса «ИИ на заказ»? Если нет — как бы вы назвали?",
    "«haftasiga 2–3 kechqurun» — так говорят про «2–3 вечера в неделю»? В рекламе вы поменяли «2 kechqurun» на «2 kunda» — здесь тоже нужно иначе?",
    "Первый экран: «Onlayn kasbni egallang: AI-mutaxassis» — естественно? Есть ли вариант живее?",
    "«Mini-kurs ustozi» (заголовок блока про Никиту) и «mini-kursni olib boruvchi» (под подписью на сертификате) — подходят?",
    "Реквизиты в подвале: «IP (yakka tartibdagi tadbirkor, Rossiya)», «Hisob-kitob raqami», «MChJ», «Korrespondent hisobvaraq» — правильно ли так писать российские ИП и банк?",
    "«aldov» для русского «лохотрон» — не слишком мягко или грубо? Как говорят ваши сверстники?",
]


def plain(fragment: str) -> str:
    text = re.sub(r"<[^>]+>", "", fragment)
    text = html.unescape(text).replace(" ", " ").replace("⁠", "").replace("‑", "-")
    return re.sub(r"\s+", " ", text).strip()


def page_rows() -> list[tuple[str, str, str, str]]:
    """(раздел, подпись, русский, узбекский) по ключам data-i18n в порядке страницы."""
    page = (ROOT / "ru" / "index.html").read_text(encoding="utf-8")
    ru: dict[str, str] = {}
    for m in re.finditer(r'<(\w+)\b[^>]*\sdata-i18n="([\w.]+)"[^>]*>', page):
        tag, key = m.group(1), m.group(2)
        if key in ru:
            continue
        depth, pos = 1, m.end()
        for t in re.finditer(rf"<(/?){tag}\b[^>]*>", page[pos:]):
            depth += -1 if t.group(1) else 1
            if depth == 0:
                ru[key] = plain(page[pos:pos + t.start()])
                break
    grouped = {k: g for g in GROUPS for k in g}
    hooks = {k.rsplit(".", 1)[0] for k in ru if k.startswith("hook.")}
    for h in hooks:
        grouped.update({f"{h}.{s}": tuple(f"{h}.{x}" for x in "abc") for s in "abc"})
    rows, done = [], set()
    for key in ru:
        if key in done:
            continue
        keys = grouped.get(key, (key,))
        done.update(keys)
        prefix = key.split(".")[0]
        ru_text = LINE.join(ru[k] for k in keys)
        uz_text = LINE.join(plain(UZ["text"][k]) for k in keys)
        if keys[-1] == "b1.lead":  # подзаголовок — отдельной фразой под заголовком
            ru_text = LINE.join(ru[k] for k in keys[:2]) + " → " + ru[keys[2]]
            uz_text = LINE.join(plain(UZ["text"][k]) for k in keys[:2]) + " → " + plain(UZ["text"][keys[2]])
        label = "заголовок «" + key.split(".")[1] + "»" if prefix == "hook" else LABELS.get(key) or _auto_label(key)
        rows.append((SECTIONS[prefix], label, ru_text, uz_text))
    return rows


def _auto_label(key: str) -> str:
    k = key.split(".")[1]
    for pat, name in ((r"l\d", "пункт списка"), (r"n\d", "номер урока"), (r"t\d", "название урока"), (r"r\d", "результат урока"),
                      (r"d\d", "описание урока"), (r"s2[a-e]", "подпись иконки"), (r"rg\db", "цифра"), (r"rg\d", "к цифре"),
                      (r"p\d", "пункт"), (r"q\d", "вопрос"), (r"a\d", "ответ")):
        if re.fullmatch(pat, k):
            return name
    return k


def hidden_rows() -> list[tuple[str, str, str, str]]:
    sec = "Скрытый текст: вкладка браузера, превью ссылки в Telegram, подписи картинок для незрячих"
    rows = []
    for ru, uz in UZ["replace"]:
        val = lambda s: plain(re.sub(r'^[\w:-]+="|"$', "", s.strip()))  # noqa: E731
        rows.append((sec, "", val(ru), val(uz)))
    return rows


def privacy_rows() -> list[tuple[str, str, str, str]]:
    sec = "Политика конфиденциальности (отдельная страница, ссылка в подвале)"
    grab = lambda p: [plain(x) for x in re.findall(r"<(?:h1|h2|p|a class=\"back\")[^>]*>(.*?)</(?:h1|h2|p|a)>", p.read_text(encoding="utf-8"), re.S)]  # noqa: E731
    ru, uz = grab(ROOT / "ru" / "privacy.html"), grab(ROOT / "privacy.html")
    if len(ru) != len(uz):
        raise SystemExit(f"политика: абзацев RU {len(ru)}, UZ {len(uz)}")
    return [(sec, "", r.replace(" · O‘zbekcha", ""), u.replace(" · Русская версия", "")) for r, u in zip(ru, uz)]


def build(rows: list[tuple[str, str, str, str]]) -> Path:
    grey, col_w = RGBColor(0x6B, 0x6B, 0x6B), (Cm(8.6), Cm(8.6), Cm(8.5))

    def shade(cell, color):
        shd = OxmlElement("w:shd")
        for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", color)):
            shd.set(qn(k), v)
        cell._tc.get_or_add_tcPr().append(shd)

    def fill(cell, text, *, size=10, bold=False, color=None):
        cell.paragraphs[0].text = ""
        run = cell.paragraphs[0].add_run(text)
        run.font.size, run.bold = Pt(size), bold
        if color:
            run.font.color.rgb = color

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    for name in ("Normal", "Title", "Heading 1"):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        if name != "Normal":
            st.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("Страница мини-курса: русский и узбекский", level=0)
    for line in INTRO:
        doc.add_paragraph(line)
    doc.add_paragraph("Что нужно:").runs[0].bold = True
    for line in ASKS:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_paragraph("Вопросы, в которых мы не уверены (ответьте прямо здесь или в правках):").runs[0].bold = True
    for line in QUESTIONS:
        doc.add_paragraph(line, style="List Number")

    table = doc.add_table(rows=1, cols=3)
    table.style, table.autofit = "Table Grid", False
    head = table.rows[0]
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    head._tr.get_or_add_trPr().append(hdr)
    for cell, title in zip(head.cells, ("Русский", "Узбекский (перевод ИИ)", "Правки")):
        fill(cell, title, bold=True)
        shade(cell, "EDEDED")
    sec_prev = None
    for sec_name, label, ru, uz in rows:
        if sec_name != sec_prev:
            merged = table.add_row().cells
            merged = merged[0].merge(merged[2])
            fill(merged, sec_name, bold=True)
            shade(merged, "F2F2F2")
            sec_prev = sec_name
        cells = table.add_row().cells
        if label:
            fill(cells[0], f"{label}: ", size=8, color=grey)
            cells[0].paragraphs[0].add_run(ru).font.size = Pt(10)
        else:
            fill(cells[0], ru)
        fill(cells[1], uz)
        fill(cells[2], "")
        shade(cells[2], "FFFDF0")
    for row in table.rows:
        for cell, w in zip(row.cells, col_w):
            cell.width = w

    doc.add_heading("Общие замечания", level=1).paragraph_format.page_break_before = True
    doc.add_paragraph("Если есть что-то общее для всего текста — напишите здесь.")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / DOCX_NAME
    doc.save(out)
    shutil.copy(out, DESK / DOCX_NAME)
    return out


def main() -> None:
    order = list(SECTIONS.values())  # первый экран, кнопка, уроки… заголовки под рекламу — в конце страницы
    rows = sorted(page_rows(), key=lambda r: order.index(r[0])) + hidden_rows() + privacy_rows()
    out = build(rows)
    print(f"docx → {out.relative_to(ROOT.parent.parent)}: {len(rows)} строк, копия на Рабочем столе")


if __name__ == "__main__":
    main()
